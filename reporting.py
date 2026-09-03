import os
import zipfile
import openpyxl
from datetime import datetime


def generate_approved_note_numbers_report(db):
    app_data = os.path.join(os.path.expanduser("~"), ".instapay_receipts")
    os.makedirs(app_data, exist_ok=True)
    path = os.path.join(app_data, "Approved_Note_Numbers.xlsx")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Approved Note Numbers"
    ws.append(["Note Number", "Corresponding Name"])
    for pair in db.get_approved_note_pairs():
        ws.append([pair["note_number"], pair["name"]])
    ws.column_dimensions["A"].width = 24
    ws.column_dimensions["B"].width = 32
    wb.save(path)
    return path


class Validator:
    @staticmethod
    def validate_receipt_data(data, db):
        errors = {}

        # 1. Receipt Number
        receipt_no = str(data.get("receipt_number", "")).strip()
        if not receipt_no:
            errors["receipt_number"] = "Receipt number is required."
        elif db.receipt_exists(receipt_no):
            errors["receipt_number"] = f"Receipt number '{receipt_no}' has already been saved."

        # 2. Account Number / Label
        account_no = str(data.get("account_number", "")).strip()
        if not account_no:
            errors["account_number"] = "Account number/label is required."
        elif not db.is_account_approved(account_no):
            errors["account_number"] = f"Account number/label '{account_no}' is not in the approved accounts list."

        # 3. Receipt Date / Time
        date_str = str(data.get("receipt_date", "")).strip()
        if not date_str:
            errors["receipt_date"] = "Receipt date and time is required."

        # 4. Note Number
        note_number = str(data.get("note_number", "")).replace(",", "").strip()
        if not note_number:
            errors["note_number"] = "A number is required in the Note section."
        elif not note_number.isdigit():
            errors["note_number"] = "Note number must contain digits only."
        elif not db.is_note_number_approved(note_number):
            errors["note_number"] = f"Note number '{note_number}' is not in the approved note numbers list."

        # 5. Money Value
        money_raw = data.get("money_value", "")
        try:
            money_val = float(str(money_raw).replace(",", ""))
            if money_val <= 0:
                errors["money_value"] = "Money value must be greater than zero."
        except ValueError:
            errors["money_value"] = "Money value must be a valid positive number."

        return errors


def _parse_receipt_date(value):
    """Return a receipt date as a date object, or None when it cannot be parsed."""
    value = str(value or "").strip()
    if not value:
        return None

    formats = (
        "%d %b %Y %I:%M %p",
        "%d %b %Y %H:%M",
        "%d-%b-%Y %I:%M %p",
        "%d-%b-%Y %H:%M",
        "%Y-%m-%d %H:%M:%S",
        "%Y-%m-%d %H:%M",
    )
    for date_format in formats:
        try:
            return datetime.strptime(value, date_format).date()
        except ValueError:
            continue

    try:
        return datetime.fromisoformat(value.replace("Z", "+00:00")).date()
    except ValueError:
        return None


def _record_was_scanned_on(r, target_date):
    created_str = r.get("created_at", "")
    try:
        return datetime.fromisoformat(created_str).date() == target_date.date()
    except (TypeError, ValueError):
        return True


def _safe_sheet_title(date_str):
    return date_str[:31]


def _format_report_date(value):
    parsed = value if hasattr(value, "year") else _parse_receipt_date(value)
    if parsed is None:
        return ""
    return f"{parsed.day}-{parsed.month}-{parsed.year}"


def _format_report_value(value):
    try:
        amount = float(str(value or "").replace(",", "").strip())
    except (TypeError, ValueError):
        return str(value or "").strip()
    if amount.is_integer():
        return f"{amount:,.0f}"
    return f"{amount:,.2f}".rstrip("0").rstrip(".")


def _write_receipt_workbook(path, date_str, receipts):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = _safe_sheet_title(date_str)

    headers = ["Date", "Value", "Client Code", "Receipt No"]
    ws.append(headers)

    for cell in ws[1]:
        cell.font = openpyxl.styles.Font(bold=True, color="FFFFFF")
        cell.fill = openpyxl.styles.PatternFill(
            start_color="1F4E78", end_color="1F4E78", fill_type="solid"
        )

    for r in receipts:
        row_date = _parse_receipt_date(r.get("receipt_date"))
        ws.append([
            _format_report_date(row_date or _parse_receipt_date(date_str)),
            _format_report_value(r.get("money_value", "")),
            r.get("note_number", ""),
            r.get("receipt_number", ""),
        ])

    widths = {"A": 16, "B": 16, "C": 30, "D": 24}
    for column, width in widths.items():
        ws.column_dimensions[column].width = width
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    wb.save(path)


def generate_daily_reports(db, target_date=None):
    """Create the current-day folder, date-specific Excel files, and one batch image ZIP.

    The folder is based on the day the receipts are scanned. Excel filenames are based
    on each receipt's own date, so older receipt dates are reported separately in the
    same current-day folder. The ZIP contains every receipt scanned in the batch.

    Returns (excel_paths, zip_path, folder_path, scan_date_str).
    """
    if not target_date:
        target_date = datetime.now()

    scan_date_str = target_date.strftime("%d-%b-%Y")
    app_data = os.path.join(os.path.expanduser("~"), ".instapay_receipts")
    daily_folder = os.path.join(app_data, scan_date_str)
    os.makedirs(daily_folder, exist_ok=True)

    all_receipts = db.get_all_receipts()
    scanned_receipts = [r for r in all_receipts if _record_was_scanned_on(r, target_date)]

    grouped = {}
    for record in scanned_receipts:
        receipt_date = _parse_receipt_date(record.get("receipt_date"))
        # Validated records should have a date. This fallback prevents data loss
        # if an older legacy record cannot be parsed.
        report_date = receipt_date or target_date.date()
        grouped.setdefault(report_date, []).append(record)

    excel_paths = []
    for report_date in sorted(grouped):
        report_date_str = report_date.strftime("%d-%b-%Y")
        excel_filename = f"Instapay_Receipts_{report_date_str}.xlsx"
        excel_path = os.path.join(daily_folder, excel_filename)
        _write_receipt_workbook(excel_path, report_date_str, grouped[report_date])
        excel_paths.append(excel_path)

    generate_approved_note_numbers_report(db)

    zip_filename = f"Instapay_Receipts_Images_{scan_date_str}.zip"
    zip_path = os.path.join(daily_folder, zip_filename)
    with zipfile.ZipFile(zip_path, "w", zipfile.ZIP_DEFLATED) as zipf:
        used_names = set()
        for r in scanned_receipts:
            img_path = r.get("image_path")
            if img_path and os.path.exists(img_path):
                ext = os.path.splitext(img_path)[1]
                base_name = f"Receipt_{r['receipt_number']}{ext}"
                arcname = base_name
                counter = 2
                while arcname in used_names:
                    arcname = f"Receipt_{r['receipt_number']}_{counter}{ext}"
                    counter += 1
                used_names.add(arcname)
                zipf.write(img_path, arcname=arcname)

    return excel_paths, zip_path, daily_folder, scan_date_str
